"""Build report/report.docx (and report/report.pdf via Microsoft Word) from outputs/results.json.

    python -m src.build_report
"""
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .data import FIG_DIR, ROOT
from .plots import pretty_feature
from .results_view import ResultsView, f3, f4, n, pct, pp

OUT_DOCX = ROOT / "report" / "report.docx"
OUT_PDF = ROOT / "report" / "report.pdf"
NAVY = RGBColor(0x14, 0x23, 0x40)
GREY = RGBColor(0x52, 0x51, 0x4E)


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self._styles()

    def _styles(self):
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        pf = normal.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.15
        for lvl, size in [(1, 16), (2, 13), (3, 11.5)]:
            h = st[f"Heading {lvl}"]
            h.font.name = "Cambria"
            h.element.rPr.rFonts.set(qn("w:asciiTheme"), "")  # override theme font
            h.element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
            h.element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.color.rgb = NAVY
            h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
        cap = st["Caption"]
        cap.font.name = "Calibri"
        cap.font.size = Pt(9.5)
        cap.font.italic = False
        cap.font.bold = False
        cap.font.color.rgb = GREY
        cap.paragraph_format.space_after = Pt(10)
        for sec in self.doc.sections:
            sec.top_margin = sec.bottom_margin = Cm(2.3)
            sec.left_margin = sec.right_margin = Cm(2.4)

    # ---------------------------------------------------------------- blocks
    def h(self, text, level=1):
        return self.doc.add_heading(text, level=level)

    def p(self, *parts, align=WD_ALIGN_PARAGRAPH.JUSTIFY, style=None, space_after=None):
        """parts: strings, or (text, 'b'|'i'|'bi'|'ph') tuples. 'ph' = highlighted placeholder."""
        para = self.doc.add_paragraph(style=style)
        para.alignment = align
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        for part in parts:
            txt, fmt = (part, "") if isinstance(part, str) else part
            r = para.add_run(txt)
            r.bold = "b" in fmt
            r.italic = "i" in fmt
            if fmt == "ph":
                r.bold = True
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        return para

    def bullets(self, items, style="List Bullet"):
        for it in items:
            parts = it if isinstance(it, tuple) else (it,)
            para = self.doc.add_paragraph(style=style)
            para.paragraph_format.space_after = Pt(3)
            for part in parts:
                txt, fmt = (part, "") if isinstance(part, str) else part
                r = para.add_run(txt)
                r.bold = "b" in fmt
                r.italic = "i" in fmt

    def eq(self, text):
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = para.add_run(text)
        r.italic = True
        r.font.name = "Cambria Math"
        return para

    def figure(self, name, caption, width_cm=15.5):
        self.fig_no += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(FIG_DIR / f"{name}.png"), width=Cm(width_cm))
        c = self.doc.add_paragraph(style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(f"Figure {self.fig_no}. ")
        r.bold = True
        c.add_run(caption)
        return self.fig_no

    def next_fig(self, k=1):
        return self.fig_no + k

    def next_tab(self, k=1):
        return self.tab_no + k

    def table(self, header, rows, caption, widths=None, font_size=9.5, bold_rows=()):
        self.tab_no += 1
        c = self.doc.add_paragraph(style="Caption")
        c.paragraph_format.keep_with_next = True
        c.paragraph_format.space_after = Pt(3)
        r = c.add_run(f"Table {self.tab_no}. ")
        r.bold = True
        c.add_run(caption)
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, htxt in enumerate(header):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(htxt)
            run.bold = True
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            _shade(cell, "142340")
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(font_size)
                run.bold = ri in bold_rows
                if ri % 2 == 1:
                    _shade(cells[i], "F1F5FB")
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                cell.paragraphs[0].paragraph_format.space_after = Pt(1)
                cell.paragraphs[0].paragraph_format.line_spacing = 1.0
                if widths:
                    cell.width = Cm(widths[i])
        _repeat_header(t.rows[0])
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return self.tab_no

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _field(paragraph, instr, placeholder=""):
    r = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    paragraph.add_run(placeholder)
    r4 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r4._r.append(e)


# ======================================================================== content
def build(R: ResultsView) -> Report:
    rp = Report()
    doc = rp.doc
    d, ex, hold = R.data, R.exp, R.hold
    base = R.mic("E0")
    geo1_gain, te_gain, fe_gain = R.mic("E2") - R.mic("E1"), R.mic("E3") - R.mic("E2"), R.mic("E4") - R.mic("E3")
    ord_mic, ord_mac = R.mic("E6") - R.mic("E4"), R.mac("E6") - R.mac("E4")
    bal_mic, bal_mac = R.mic("E5") - R.mic("E4"), R.mac("E5") - R.mac("E4")
    H_mic, H_mac = hold["f1_micro"], hold["f1_macro"]
    ss = R.eda["superstructure"]
    mms, rce, cmb = (ss["has_superstructure_mud_mortar_stone"], ss["has_superstructure_rc_engineered"],
                     ss["has_superstructure_cement_mortar_brick"])
    fnd, roof = R.eda["foundation_type"], R.eda["roof_type"]
    g1 = R.eda["geo1_grade3_share"]
    kept = R.kept
    kept_txt = ", ".join(f"`{k}`".replace("`", "") for k in kept) if kept else "none of the candidates"
    best_desc = {"E4": "the multinomial model with engineered features (E4)",
                 "E5": "the class-balanced multinomial model (E5)",
                 "E6": "the ordinal model (E6)"}[R.best]

    # ------------------------------------------------------------ cover page
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Predicting Earthquake Damage to Prioritise Building Inspections")
    r.bold, r.font.size, r.font.name, r.font.color.rgb = True, Pt(24), "Cambria", NAVY
    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("A logistic-regression study of the 2015 Gorkha earthquake building-damage data (Richter's Predictor)")
    r.font.size, r.font.color.rgb = Pt(13), GREY
    for _ in range(2):
        doc.add_paragraph()
    rp.p(("Course: ", "b"), ("[COURSE CODE] – [COURSE NAME]", "ph"), align=WD_ALIGN_PARAGRAPH.CENTER)
    rp.p(("Group: ", "b"), ("[GROUP NUMBER]", "ph"), align=WD_ALIGN_PARAGRAPH.CENTER)
    rp.p(("Submission date: ", "b"), ("[DD Month YYYY]", "ph"), align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()
    mt = doc.add_table(rows=1, cols=3)
    mt.style = "Table Grid"
    mt.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htxt in enumerate(["#", "Member name", "Student ID"]):
        c = mt.rows[0].cells[i]
        c.text = ""
        rr = c.paragraphs[0].add_run(htxt)
        rr.bold = True
        rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "142340")
    for k in range(1, 7):
        cells = mt.add_row().cells
        cells[0].text = str(k)
        for j, ph in [(1, f"[MEMBER {k} NAME]"), (2, f"[STUDENT ID {k}]")]:
            cells[j].text = ""
            rr = cells[j].paragraphs[0].add_run(ph)
            rr.font.highlight_color = WD_COLOR_INDEX.YELLOW
    for row in mt.rows:
        for i, w in enumerate([1.2, 8.0, 5.0]):
            row.cells[i].width = Cm(w)
    doc.add_paragraph()
    rp.p(("Placeholders highlighted in yellow must be completed by the group before submission.", "i"),
         align=WD_ALIGN_PARAGRAPH.CENTER)
    rp.page_break()

    # ------------------------------------------------------------ abstract + TOC
    rp.h("Abstract", 1)
    rp.p(f"After the 2015 Gorkha earthquake, Nepal surveyed hundreds of thousands of damaged buildings. "
         f"We ask whether a building's recorded structure and location can predict its damage grade well enough to rank buildings for inspection. "
         f"Using {n(d['n_train_rows'])} labelled buildings from the DrivenData competition Richter's Predictor, we built a leakage-safe scikit-learn pipeline "
         f"and ran seven controlled experiments, restricted to logistic regression. "
         f"Location, encoded with cross-fitted multiclass target encoding of {n(d['n_unique']['geo_level_3_id'])} local areas, gave the largest single gain "
         f"({pp(te_gain)} in cross-validated micro-F1). An ordinal formulation using two cumulative binary models "
         f"{'did not outperform' if ord_mic <= 0.001 else 'outperformed'} the standard multinomial model, and class weighting traded overall accuracy for recall of the rare low-damage class. "
         f"The tuned model ({R.best}, C = {R.best_C:g}) reaches a holdout micro-F1 of {f3(H_mic)} and macro-F1 of {f3(H_mac)}, against {f3(base)} for the majority-class baseline. "
         f"Ranking buildings by predicted probability of grade-3 damage, inspecting the top 20% of buildings finds {pct(R.cap(20))} of all grade-3 buildings, "
         f"{R.lift(20):.1f} times what a random order finds. The coefficients agree with engineering intuition: "
         f"mud-mortar stone and stone-flag construction raise the odds of destruction, and engineered reinforced concrete and cement-bonded masonry lower them.")
    rp.h("Contents", 1)
    toc = doc.add_paragraph()
    _field(toc, 'TOC \\o "1-2" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    rp.page_break()

    # ------------------------------------------------------------ 1 intro
    rp.h("1. Introduction and Problem Statement", 1)
    rp.p("On 25 April 2015 a magnitude-7.8 earthquake struck the Gorkha district of Nepal. It killed nearly 9,000 people and damaged or destroyed "
         "hundreds of thousands of buildings (National Planning Commission, 2015). After a disaster on this scale, trained structural inspectors are a scarce resource. "
         "Every building must eventually be assessed, but the order of inspection decides how quickly dangerous structures are cordoned off, "
         "families are rehoused, and reconstruction aid is targeted.")
    rp.p("The Government of Nepal later released one of the largest post-disaster building surveys ever collected. A version of it underlies the DrivenData "
         "competition ", ("Richter's Predictor: Modeling Earthquake Damage", "i"), " (DrivenData, 2019). This project uses those data to answer one question:")
    q = rp.p(("“Can we predict how badly a building was damaged in the 2015 Gorkha earthquake from its structure and location, "
              "so post-disaster inspectors can prioritise which buildings to assess first?”", "i"), align=WD_ALIGN_PARAGRAPH.CENTER)
    q.paragraph_format.left_indent = q.paragraph_format.right_indent = Cm(1.2)
    rp.p("We treat this as a three-class classification problem with an ordered target: damage grade 1 (low), 2 (medium) and 3 (almost complete destruction). "
         "We deliberately restrict ourselves to ", ("logistic regression", "b"),
         ". Its predictions are fast to compute, its coefficients can be audited by engineers, and its probabilities can be used directly to rank buildings. "
         "Our contributions are:")
    rp.bullets([
        "a fully leakage-safe preprocessing and modelling pipeline, evaluated with stratified 5-fold cross-validation and a holdout that is used only once;",
        "a controlled sequence of seven experiments that isolates the value of structure, coarse location, fine location, engineered features, class weighting and an ordinal formulation;",
        "cross-fitted multiclass target encoding of the high-cardinality location identifiers, which turns out to be the single most valuable idea;",
        "a prioritisation analysis that converts model output into an inspection policy, and a coefficient-level interpretation in plain language.",
    ])

    # ------------------------------------------------------------ 2 dataset
    rp.h("2. Dataset Description", 1)
    rp.p(f"The labelled data consist of two files joined on ", ("building_id", "i"), f": ", ("train_values.csv", "i"),
         f" with {n(d['n_train_rows'])} rows and {d['n_features']} features, and ", ("train_labels.csv", "i"),
         f" with the damage grade. A further {n(d['n_test_rows'])} buildings in ", ("test_values.csv", "i"),
         " have no labels; we use them only to produce a competition submission. Table "
         f"{rp.next_tab()} groups the features.")
    rp.table(["Group", "Features", "Type"], [
        ["Location", f"geo_level_1_id ({d['n_unique']['geo_level_1_id']} regions), geo_level_2_id ({n(d['n_unique']['geo_level_2_id'])}), "
                     f"geo_level_3_id ({n(d['n_unique']['geo_level_3_id'])})", "nested categorical ids"],
        ["Size & age", "age, area_percentage, height_percentage, count_floors_pre_eq", "numeric (normalised)"],
        ["Construction", "foundation_type, roof_type, ground_floor_type, other_floor_type, position, plan_configuration, land_surface_condition",
         "categorical (anonymised letters)"],
        ["Superstructure", "11 has_superstructure_* flags (adobe mud, mud-mortar stone, stone flag, cement-mortar stone/brick, "
                           "mud-mortar brick, timber, bamboo, RC non-engineered, RC engineered, other)", "binary"],
        ["Use & ownership", "legal_ownership_status, count_families, 11 has_secondary_use_* flags", "categorical / numeric / binary"],
    ], "Feature groups in the Richter's Predictor data.", widths=[3.0, 9.5, 3.8])
    rp.p(f"We verified the facts given in the brief rather than assuming them, and the notebook asserts each one. There are {d['n_missing_train']} missing values in "
         f"either file, {d['n_categorical']} categorical columns and {d['n_binary']} binary columns. The class distribution is imbalanced: "
         f"grade 1 {pct(d['class_share']['1'])}, grade 2 {pct(d['class_share']['2'])} and grade 3 {pct(d['class_share']['3'])} (Figure {rp.next_fig()}). "
         f"Two data-quality issues matter. First, {n(d['n_age_placeholder'])} buildings ({pct(d['share_age_placeholder'], 2)}) have ", ("age = 995", "i"),
         f", a placeholder far beyond the largest real age of {d['age_max_real']} years. Second, the finest location level has "
         f"{n(d['n_unique']['geo_level_3_id'])} distinct values, too many to one-hot encode sensibly. "
         f"Also, {pct(d['test_geo3_unseen_share'], 2)} of test buildings fall in a geo_level_3 area that never appears in the labelled data.")
    rp.figure("fig01_class_distribution", "Damage-grade distribution in the 80% training split. Grade 2 is the majority class; grade 1 is rare.", 11)

    # ------------------------------------------------------------ 3 preprocessing
    rp.h("3. Data Pre-processing", 1)
    rp.p(f"Before any exploration, we made a stratified 80/20 split ({n(R.split['n_train'])} training and {n(R.split['n_holdout'])} holdout buildings). "
         "All exploration below uses only the training split, so nothing we decided about the features was influenced by the holdout.")
    rp.h("3.1 Exploration", 2)
    rp.p(f"Construction type separates the classes sharply (Figure {rp.next_fig()}). Foundation type r, used by {pct(fnd['prevalence']['r'], 0)} of buildings, "
         f"has a grade-3 share of {pct(fnd['grade3_share']['r'], 0)}. Foundation i has a grade-3 share of only {pct(fnd['grade3_share']['i'], 0)} and a grade-1 share of {pct(fnd['grade1_share']['i'], 0)}. "
         f"Roof type x shows the same protective pattern ({pct(roof['grade3_share']['x'], 0)} grade 3). The category codes are anonymised, but these profiles are "
         "consistent with engineered, reinforced construction.")
    fig_found = rp.figure("fig02_damage_by_foundation_roof", "Damage-grade mix by foundation type (left) and roof type (right), training split. Bars are sorted by grade-3 share; n is the number of buildings.")
    rp.p(f"Superstructure material is the clearest structural signal (Figure {rp.next_fig()}). Mud-mortar stone is used in {pct(mms['prevalence'], 0)} of buildings, "
         f"and {pct(mms['grade3_share'], 0)} of those reached grade 3. Engineered reinforced concrete (RC) reached grade 3 in only {pct(rce['grade3_share'], 1)} of cases, "
         f"with {pct(rce['grade1_share'], 0)} at grade 1. Cement-mortar brick is similarly protective ({pct(cmb['grade3_share'], 0)} grade 3).")
    rp.figure("fig03_damage_by_superstructure", "Damage-grade mix by superstructure material. A building can use several materials, so rows overlap.", 14)
    rp.p(f"Location matters at least as much as structure (Figure {rp.next_fig()}). Across the {d['n_unique']['geo_level_1_id']} top-level regions, the grade-3 share ranges from "
         f"{pct(g1['min'], 0)} (region {g1['argmin']}) to {pct(g1['max'], 0)} (region {g1['argmax']}). Location most likely stands in for shaking intensity, "
         "that is, distance from the epicentre and local ground conditions. These are not in the data, but they strongly shape damage.")
    rp.figure("fig04_damage_by_geo1", f"Damage-grade mix for each of the {d['n_unique']['geo_level_1_id']} geo_level_1 regions, sorted by grade-3 share.")
    sk, skl = R.eda["skewness"], R.eda["skewness_after_log1p"]
    rp.p(f"The numeric features are right-skewed, with skewness {sk['age']:.2f} for age, {sk['area_percentage']:.2f} for area and {sk['height_percentage']:.2f} for height. "
         f"A log1p transform reduces this to {skl['age']:.2f}, {skl['area_percentage']:.2f} and {skl['height_percentage']:.2f} (Appendix C, Figure C1). "
         f"Individually these features are only weakly related to damage: the largest absolute Spearman correlation with damage grade among numeric and material features is "
         f"{max(abs(v) for v in R.eda['spearman_with_target'].values()):.2f}. Height and number of floors are strongly collinear "
         f"(ρ = {R.eda['spearman_height_floors']:.2f}; Appendix C, Figure C2). The weak marginal correlations show that damage comes from combinations of many weak signals, "
         "which is what a multivariate model needs to capture. The L2 penalty of logistic regression keeps the collinear pair stable.")

    rp.h("3.2 Cleaning", 2)
    rp.p(f"No imputation of genuinely missing values was needed. The age placeholder was handled explicitly. Buildings with age = 995 have a lower grade-3 share "
         f"({pct(R.eda['age_placeholder_grade3_share'])}) than the rest ({pct(R.eda['age_real_grade3_share'])}), so the placeholder carries information and should not be "
         "treated as a real age of 995 years. A custom transformer (AgeCleaner) therefore (i) adds a binary indicator age_is_placeholder and (ii) replaces the "
         "placeholder with the median age of the training fold. Because the median is learned in fit, it is recomputed inside every cross-validation fold.")

    rp.h("3.3 Encoding and scaling", 2)
    rp.p(f"All transformations are implemented as one scikit-learn ColumnTransformer inside a Pipeline (Table {rp.next_tab()}), so every learned statistic comes from training data only.")
    rp.table(["Columns", "Transformation", "Rationale"], [
        ["age, area_percentage, height_percentage", "log1p, then StandardScaler", "reduce right skew; common scale for L2 and lbfgs"],
        ["count_floors_pre_eq, count_families (+ engineered)", "StandardScaler", "small counts, little skew"],
        ["8 categorical columns", "OneHotEncoder (all levels, handle_unknown='ignore')", "nominal codes; L2 keeps full dummy sets identifiable"],
        ["geo_level_1_id (E2+)", "OneHotEncoder", f"{d['n_unique']['geo_level_1_id']} regions, small enough to one-hot"],
        ["geo_level_2_id, geo_level_3_id (E3+)", "TargetEncoder(target_type='multiclass'), 5-fold cross-fitting", "high cardinality; 3 smoothed class rates per level"],
        ["22 binary flags + age_is_placeholder", "passthrough", "already 0/1"],
    ], "Preprocessing pipeline. Every step is fitted on training folds only.", widths=[5.0, 5.8, 5.5])
    rp.p("The multiclass TargetEncoder replaces each geo id by three numbers: the smoothed proportions of grade-1, grade-2 and grade-3 buildings observed for that id. "
         "Rare ids are shrunk toward the global class distribution, and ids never seen in training receive the global distribution. Target encoding can leak the label, "
         "because a building's own grade would contribute to its own feature. scikit-learn's implementation therefore uses ", ("cross-fitting", "i"),
         " in fit_transform: training rows are encoded with statistics computed on the other internal folds (Micci-Barreca, 2001; Pedregosa et al., 2011).")

    rp.h("3.4 Feature engineering", 2)
    abl = R.abl["candidates"]
    rp.p(f"Four candidate features were constructed: n_superstructure (number of materials used), log_area_height_ratio (footprint relative to height, a slenderness proxy), "
         f"floors_x_log_age (taller and older buildings), and height_per_floor. Fold-to-fold noise is of the same order as the expected effects, so each candidate was "
         f"added to E3 on its own and compared on the same five folds. A candidate was kept only if its mean paired gain was positive ", ("and", "i"),
         f" it improved at least four of the five folds. Table {rp.next_tab()} shows the outcome; the features kept were {kept_txt}.")
    rp.table(["Candidate", "CV micro-F1", "Mean paired gain", "Folds improved", "Decision"],
             [[f, f4(a["f1_micro_mean"]), pp(a["mean_gain"], 2), f"{a['folds_improved']}/5", "kept" if a["kept"] else "dropped"]
              for f, a in abl.items()],
             f"Add-one ablation of engineered features on top of E3 (E3 micro-F1 = {f4(R.mic('E3'))}).", widths=[4.5, 2.6, 3.2, 2.8, 2.4])

    # ------------------------------------------------------------ 4 methodology
    rp.h("4. Methodology", 1)
    rp.h("4.1 Training and testing procedure", 2)
    rp.bullets([
        (("Split. ", "b"), f"values and labels were merged on building_id and split 80/20, stratified by damage grade, with random_state = 42."),
        (("Model selection. ", "b"), "every experiment was scored by stratified 5-fold cross-validation (shuffled, random_state = 42) on the 80% split. "
                                     "We report the mean and standard deviation over folds. Identical folds across experiments allow paired comparisons."),
        (("Tuning. ", "b"), f"the best of E4–E6 by mean CV micro-F1 was tuned over C ∈ {{{', '.join(f'{c:g}' for c in R.tun['grid'])}}} "
                            "with GridSearchCV, using the same folds and scoring f1_micro. GridSearchCV then refitted the best setting on the whole 80% split."),
        (("Final test. ", "b"), "the refitted model was scored once on the 20% holdout; the notebook asserts that this happens only once."),
        (("Submission. ", "b"), f"the tuned configuration was refitted on all {n(R.sub['refit_rows'])} labelled buildings to predict the {n(R.sub['n_rows'])} test buildings."),
    ])
    rp.h("4.2 Models", 2)
    rp.p("The multinomial logistic regression (Hosmer et al., 2013; Hastie et al., 2009) models the class probabilities with a softmax over linear scores:")
    rp.eq("P(y = k | x) = exp(β₀ₖ + βₖᵀx) / Σⱼ exp(β₀ⱼ + βⱼᵀx),   k ∈ {1, 2, 3}")
    rp.p("It is fitted by minimising the cross-entropy plus an L2 penalty ‖β‖²/(2C) with the lbfgs solver. "
         f"All inputs are scaled or binary, and we allowed up to {R.conv['max_iter_allowed']} iterations. The largest number of iterations used by any fit "
         f"(experiments, ablation, all {len(R.tun['grid'])}×5 grid fits and the refits) was {R.conv['max_n_iter_observed']}. "
         "Every fit therefore converged, and no ConvergenceWarning was raised.")
    rp.p("For the ordinal model (E6) we follow Frank and Hall (2001). We fit two binary logistic models on the same features, for P(y > 1 | x) and P(y > 2 | x). "
         "Class probabilities are differences of the cumulative probabilities:")
    rp.eq("P(y=1) = 1 − P(y>1),   P(y=2) = P(y>1) − P(y>2),   P(y=3) = P(y>2)")
    rp.p("The two models are fitted independently, so their cumulative probabilities can occasionally cross. We enforce P(y>2) ≤ P(y>1) before taking differences, "
         "and we predict the class with the highest probability. Unlike the proportional-odds model (McCullagh, 1980), the two models do not share slopes.")
    rp.h("4.3 Experimental design", 2)
    rp.p(f"Each experiment changes one thing relative to its predecessor (Table {rp.next_tab()}).")
    rp.table(["ID", "Configuration", "Question answered"], [
        ["E0", "Majority-class baseline (always grade 2)", "What does guessing achieve?"],
        ["E1", "LR, all non-location features", "How much does the building itself tell us?"],
        ["E2", "E1 + one-hot geo_level_1_id", "Value of coarse location"],
        ["E3", "E2 + target-encoded geo_level_2_id and geo_level_3_id", "Value of fine location"],
        ["E4", f"E3 + engineered features ({', '.join(kept) if kept else 'none kept'})", "Value of feature engineering"],
        ["E5", "E4 with class_weight='balanced'", "Does reweighting help the rare class?"],
        ["E6", "E4 as ordinal (two cumulative binary LRs)", "Does exploiting the order help?"],
    ], "Experiment design. All logistic models use C = 1 before tuning.", widths=[1.2, 8.0, 7.0])
    rp.h("4.4 Evaluation metrics", 2)
    rp.p("Micro-averaged F1 is the competition metric. For single-label multiclass prediction it equals accuracy. Macro-averaged F1 gives each grade equal weight and "
         "so shows performance on the rare grade 1. On the holdout we also report per-class precision and recall and a row-normalised confusion matrix. "
         "For the prioritisation analysis, the capture rate at x% is the share of all grade-3 buildings found among the x% of buildings with the highest predicted P(grade 3). "
         "Its ratio to x is the lift.")

    # ------------------------------------------------------------ 5 results
    rp.h("5. Experimental Study and Analysis", 1)
    rp.h("5.1 Cross-validation results", 2)
    rows = [[k, v["description"], R.cv_cell(k), R.cv_cell(k, "f1_macro"), f"{R.rec(k, 1):.3f}"] for k, v in ex.items()]
    rows.append([f"{R.best}*", f"{ex[R.best]['description']}, tuned C = {R.best_C:g}",
                 f"{R.tun['best_f1_micro_mean']:.4f} ± {R.tun['best_f1_micro_std']:.4f}", f"{R.tun['best_f1_macro_mean']:.4f}", "–"])
    rp.p(f"Table {rp.next_tab()} and Figure {rp.next_fig()} summarise the experiments.")
    rp.table(["ID", "Description", "Micro-F1 (mean ± sd)", "Macro-F1 (mean ± sd)", "Grade-1 recall"], rows,
             "Stratified 5-fold CV results on the 80% training split. * = tuned final model.",
             widths=[1.2, 6.3, 3.3, 3.3, 2.2], font_size=9, bold_rows=(len(rows) - 1,))
    rp.figure("fig07_experiments", "CV micro-F1 by experiment (error bars: ±1 sd over folds). Orange marks the configuration selected for tuning.")
    rp.p(f"Structure alone (E1) improves on the {f3(base)} baseline to {f3(R.mic('E1'))}. That is a real gain, but most buildings still end up in the majority class. "
         f"Coarse location (E2) adds {pp(geo1_gain)}, and target-encoded fine location (E3) adds a further {pp(te_gain)}, the largest single step in the study. "
         f"The macro-F1 rises from {f3(R.mac('E1'))} to {f3(R.mac('E3'))} along the same path. Engineered features (E4) change micro-F1 by only {pp(fe_gain, 2)}. "
         "The raw features already contain most of what a linear model can extract, and the remaining signal is in location rather than geometry. "
         f"Fold standard deviations are at most {max(R.mic_sd(k) for k in ex):.4f}, so differences of more than about half a percentage point are robust.")
    rp.p((f"Class weighting (E5) lowers micro-F1 by {pp(-bal_mic, 2).lstrip('+')} and changes macro-F1 by {pp(bal_mac, 2)}. " if bal_mic < 0 else
          f"Class weighting (E5) changes micro-F1 by {pp(bal_mic, 2)} and macro-F1 by {pp(bal_mac, 2)}. ") +
         f"Grade-1 recall rises from {R.rec('E4', 1):.3f} to {R.rec('E5', 1):.3f}, and grade-2 recall falls from {R.rec('E4', 2):.3f} to {R.rec('E5', 2):.3f}. "
         f"The ordinal model (E6) scores {f4(R.mic('E6'))} micro-F1 against {f4(R.mic('E4'))} for the multinomial model. Section 6.1 discusses both. "
         f"Since micro-F1 is the competition metric and the tuning criterion, {best_desc} was carried forward to tuning.")
    rp.h("5.2 Hyperparameter tuning", 2)
    trows = R.tun["rows"]
    lo = min(trows, key=lambda r: r["C"])
    spread = max(r["mean"] for r in trows) - min(r["mean"] for r in trows)
    rp.p(f"GridSearchCV selected C = {R.best_C:g} with CV micro-F1 {f4(R.tun['best_f1_micro_mean'])} ± {f4(R.tun['best_f1_micro_std'])}. The curve is essentially flat "
         f"(Appendix B, Figure B1): across four orders of magnitude of C, mean micro-F1 varies by only {spread * 100:.2f} percentage points, less than one fold standard deviation. "
         f"With {n(R.split['n_train'])} training rows and only {R.coef['n_features']} model features, the L2 penalty barely matters, so the choice of C is not critical. Strong regularisation "
         f"(C = {lo['C']:g}) mainly costs macro-F1 ({f4(lo['f1_macro_mean'])} vs {f4(R.tun['best_f1_macro_mean'])}), because it shrinks the coefficients that separate the rare grade 1.")
    rp.h("5.3 Holdout evaluation", 2)
    rp.p(f"On the untouched holdout ({n(hold['n'])} buildings), the tuned model achieves ", (f"micro-F1 = {f4(H_mic)}", "b"), " and ", (f"macro-F1 = {f4(H_mac)}", "b"),
         f". These are within {abs(H_mic - R.tun['best_f1_micro_mean']) * 100:.2f} percentage points of the CV estimate, which confirms that model selection did not overfit. "
         f"Per-class results are in Table {rp.next_tab()} and the confusion matrix in Figure {rp.next_fig()}.")
    rp.table(["Grade", "Precision", "Recall", "F1", "Support"],
             [[f"{c}", f3(R.pc(c, 'precision')), f3(R.pc(c, 'recall')), f3(R.pc(c, 'f1')), n(R.pc(c, 'support'))] for c in (1, 2, 3)],
             "Per-class holdout metrics of the final model.", widths=[2.0, 2.6, 2.6, 2.6, 2.6])
    rp.figure("fig09_confusion_matrix", "Row-normalised holdout confusion matrix (counts in brackets).", 10.5)
    rp.p(f"Grade 2 is recalled best ({pct(R.pc(2, 'recall'))}), then grade 3 ({pct(R.pc(3, 'recall'))}). Grade 1 is hardest ({pct(R.pc(1, 'recall'))}), "
         f"and {pct(R.cmn(1, 2))} of grade-1 buildings are predicted as grade 2. Almost all errors are between adjacent grades. Only {pct(R.cmn(3, 1), 2)} of grade-3 buildings are "
         f"predicted as grade 1, and only {pct(R.cmn(1, 3), 2)} of grade-1 buildings as grade 3. The model therefore rarely makes the costly mistake of calling a destroyed building safe.")
    rp.h("5.4 Prioritisation analysis", 2)
    rp.p(f"To connect the model to the problem statement, we ranked holdout buildings by predicted P(grade 3). We then measured the share of true grade-3 buildings found as "
         f"inspectors work down the list (Figure {rp.next_fig()}, Table {rp.next_tab()}).")
    rp.figure("fig10_prioritisation_curve", "Cumulative share of grade-3 buildings found against the share of buildings inspected, holdout. The dashed line is a random inspection order.", 12.5)
    rp.table(["Top x% inspected", "Buildings inspected", "Grade-3 captured", "Lift vs random", "Grade-3 share within list"],
             [[f"{k}%", n(R.prio['at_fraction'][str(k)]['n_inspected']), pct(R.cap(k)), f"{R.lift(k):.2f}×", pct(R.prec_top(k))] for k in (10, 20, 30)],
             f"Capture rates on the holdout ({n(R.prio['n_grade3'])} grade-3 buildings, base rate {pct(R.prio['base_rate_grade3'])}).",
             widths=[3.0, 3.2, 3.2, 3.0, 3.8])
    rp.p(f"Inspecting the top 10% of buildings finds {pct(R.cap(10))} of all grade-3 buildings, {R.lift(10):.1f} times what a random order finds. "
         f"Within that top decile, {pct(R.prec_top(10))} of buildings are actually grade 3, against a base rate of {pct(R.prio['base_rate_grade3'])}. "
         f"At 30% the capture rate is {pct(R.cap(30))}. The ROC-AUC for grade 3 against the other grades is {R.prio['auc_grade3_vs_rest']:.3f}. "
         "For an inspection agency, this means the first teams in the field would spend most of their time at buildings that really are destroyed.")
    rp.h("5.5 Interpretation of coefficients", 2)
    src = R.coef["source"]
    up, down = R.coef["top_raising"], R.coef["top_lowering"]
    mat = R.coef["materials"]
    rare = R.coef["excluded_rare_top"]
    rp.p(f"Figure {rp.next_fig()} shows the 15 features that most raise and most lower the odds of grade-3 damage in the final model ({src}). "
         "Each is shown as an odds ratio exp(β): per one standard deviation for scaled numeric features, and for presence versus absence for flags and dummy variables. "
         "In a softmax model, a dummy's coefficient is relative to the L2-centred average level, so the bars are best read as directions and relative sizes, not exact causal effects. "
         f"Two groups are excluded from the figure. Location features are discussed below. The other group is {R.coef['n_excluded_rare']} flags and dummies present in fewer than "
         f"{pct(R.coef['min_prevalence'], 0)} of buildings, whose coefficients rest on very few buildings and are unstable. For example, {pretty_feature(rare[0]['feature'])} "
         f"has an odds ratio of {rare[0]['odds_ratio']:.2f} but applies to only {pct(rare[0]['prevalence'], 2)} of buildings.")
    rp.figure("fig11_coefficients", "Top 15 risk-raising (red) and risk-lowering (blue) features for grade-3 damage, as odds ratios on a log scale. Location features and flags present in fewer than 1% of buildings are excluded.", 14.5)
    rp.p("The strongest risk-raising features are ", ", ".join(f"{pretty_feature(x['feature'])} (OR {x['odds_ratio']:.2f})" for x in up[:4]),
         ". The strongest protective features are ", ", ".join(f"{pretty_feature(x['feature'])} (OR {x['odds_ratio']:.2f})" for x in down[:4]), ". "
         "For superstructure materials, holding everything else fixed, the odds ratios are: mud-mortar stone "
         f"{mat['has_superstructure_mud_mortar_stone']['odds_ratio']:.2f}, stone flag {mat['has_superstructure_stone_flag']['odds_ratio']:.2f}, "
         f"adobe mud {mat['has_superstructure_adobe_mud']['odds_ratio']:.2f}, timber {mat['has_superstructure_timber']['odds_ratio']:.2f}, "
         f"cement-mortar brick {mat['has_superstructure_cement_mortar_brick']['odds_ratio']:.2f} and engineered RC {mat['has_superstructure_rc_engineered']['odds_ratio']:.2f}. "
         "In plain language, buildings with heavy, brittle walls of stone or mud bonded only by mud mortar break apart under strong shaking. Engineered reinforced-concrete "
         "frames and cement-bonded masonry hold together and are much more likely to escape with light damage. Older buildings are also at higher risk. The anonymised foundation and roof codes "
         f"that the model treats as protective (foundation i, roof x) are the same ones that the exploration (Figure {fig_found}) linked to low damage. That the model and the raw data agree, and that both "
         "match engineering experience, gives some confidence that the model has learned real effects rather than artefacts. The engineered-RC odds ratio is smaller in "
         "magnitude than its raw grade-3 rate would suggest, because foundation, roof and floor types that come with RC construction already carry part of that effect.")
    loc = R.coef["location_top"]
    rp.p("Location features, which are excluded from the figure, have the largest coefficients of all (for example ",
         ", ".join(f"{x['feature']} with OR {x['odds_ratio']:.2f}" for x in loc[:2]),
         "). They describe where the ground shook hardest, not what makes a building vulnerable, so they matter for prediction but give an engineer nothing to act on.")

    # ------------------------------------------------------------ 6 highlights
    rp.h("6. Highlights", 1)
    rp.h("6.1 Ordinal versus multinomial logistic regression", 2)
    rp.figure("fig12_recall_comparison", "Per-class CV recall for the multinomial (E4), class-balanced (E5) and ordinal (E6) models.", 13)
    if ord_mic <= 0.001:
        rp.p(f"Exploiting the ordering of damage grades seems natural, but on these data it ", ("did not help", "b"),
             f": the ordinal model changed micro-F1 by {pp(ord_mic, 2)} and macro-F1 by {pp(ord_mac, 2)} compared with the multinomial model on identical features "
             f"(Figure {rp.fig_no}). We see three reasons for this negative result. "
             "(i) With only three classes, the multinomial model has just one more set of coefficients than the ordinal model and can learn an ordered structure by itself. "
             "Its fitted probabilities already put very little mass on the non-adjacent grade (Section 5.3). "
             "(ii) Our ordinal model fits the two thresholds independently, so it gains nothing from sharing slopes. It can also produce crossing cumulative probabilities, which we have to clip. "
             "(iii) Ordinal structure mostly helps by avoiding large errors (1 ↔ 3), but micro-F1 counts every error the same, and large errors were already rare. "
             "This is still a useful finding: for this problem the simpler multinomial model is sufficient.")
    else:
        rp.p(f"Exploiting the ordering of damage grades helped: the ordinal model changed micro-F1 by {pp(ord_mic, 2)} and macro-F1 by {pp(ord_mac, 2)} "
             f"compared with the multinomial model on identical features (Figure {rp.fig_no}). Splitting the problem into 'damaged beyond grade 1' and "
             "'damaged beyond grade 2' lets each binary model specialise on one boundary.")
    rp.p(f"Class weighting (E5) did what it is designed to do for recall. Reweighting the rare grade 1 raised its recall from {R.rec('E4', 1):.2f} to {R.rec('E5', 1):.2f} "
         f"and grade-3 recall from {R.rec('E4', 3):.2f} to {R.rec('E5', 3):.2f}. The cost was grade-2 recall, which fell from {R.rec('E4', 2):.2f} to {R.rec('E5', 2):.2f}. "
         f"Micro-F1 therefore fell ({pp(bal_mic, 2)}), and macro-F1 {'also fell' if bal_mac < 0 else 'rose'} ({pp(bal_mac, 2)}). "
         + ("Macro-F1 falling is less obvious than micro-F1 falling. The extra grade-1 predictions are taken mostly from the large grade-2 class, so grade-1 precision drops "
            "faster than its recall rises, and grade-2 F1 falls as well. " if bal_mac < 0 else "")
         + "Whether that trade is worth making is a policy decision rather than a modelling one. An agency that cares about quickly clearing low-damage buildings, "
         "or about flagging more grade-3 buildings when acting on predicted labels rather than on a ranking, might prefer it. For the competition metric, the unweighted model is better.")
    rp.h("6.2 Cross-fitted target encoding of location", 2)
    rp.p(f"The finest location level has {n(d['n_unique']['geo_level_3_id'])} values, and level 2 has {n(d['n_unique']['geo_level_2_id'])}. One-hot encoding them would add "
         f"{n(d['n_unique']['geo_level_2_id'] + d['n_unique']['geo_level_3_id'])} sparse columns, most with only a handful of buildings, and an L2-penalised linear model would either overfit them or shrink them to nothing. "
         "Multiclass target encoding compresses each level into three smoothed damage rates, which directly answer the question 'how badly was this neighbourhood hit?'. "
         f"It produced the largest gain in the study, {pp(te_gain)} micro-F1 and {pp(R.mac('E3') - R.mac('E2'))} macro-F1 over E2. "
         "Two safeguards keep this honest. First, cross-fitting inside fit_transform means that no training building's encoding uses its own label. Second, the encoder "
         "is part of the pipeline and so is refitted inside every outer CV fold and inside GridSearchCV. The close agreement between CV and holdout scores "
         f"({f4(R.tun['best_f1_micro_mean'])} vs {f4(H_mic)}) is consistent with there being no leakage.")

    # ------------------------------------------------------------ 7 limitations
    rp.h("7. Limitations", 1)
    rp.bullets([
        (("Linear decision boundaries. ", "b"), "Logistic regression cannot represent interactions it is not given, such as material × region. Tree ensembles are known to do better "
         "on this competition, and we excluded them by design."),
        (("Location as a proxy. ", "b"), "Region ids stand in for shaking intensity, soil and distance to the fault. The model would not transfer to a different earthquake, "
         f"and {pct(d['test_geo3_unseen_share'], 2)} of test buildings fall in unseen local areas and receive only the global prior."),
        (("Labels and features recorded after the event. ", "b"), "Damage grades come from survey teams and may be noisy near grade boundaries. Some attributes, such as floor count before the earthquake, "
         "were reconstructed after it."),
        (("Anonymised categories. ", "b"), "Letter codes limit how far the coefficients can be explained in engineering terms."),
        (("Associational, not causal. ", "b"), "Odds ratios describe associations conditional on the other features. Collinear inputs (height and floors) share credit, and the "
         "softmax parameterisation makes dummy coefficients relative to an average level."),
        (("Single split. ", "b"), "The holdout estimate comes from one 20% split. CV standard deviations are small, but no confidence interval for the holdout score was computed."),
    ])

    # ------------------------------------------------------------ 8 future
    rp.h("8. Future Directions", 1)
    rp.bullets([
        (("Physical hazard features. ", "b"), "Join measured peak ground acceleration (USGS ShakeMap) and distance to the rupture, replacing region ids with physics that transfers to future events."),
        (("Interactions and splines. ", "b"), "Add material × region, floors × age and spline terms for age and area, which keeps the model linear in its parameters and interpretable."),
        (("Proper cumulative-link model. ", "b"), "Fit a proportional-odds model with shared slopes (for example statsmodels OrderedModel) and test the proportional-odds assumption."),
        (("Cost-sensitive thresholds and calibration. ", "b"), "Choose decision thresholds from inspection capacity and the relative cost of missing a destroyed building, and calibrate probabilities (Platt or isotonic)."),
        (("Uncertainty and validation. ", "b"), "Bootstrap confidence intervals for holdout metrics, and spatially grouped CV (by geo_level_2) to estimate performance in unseen areas."),
        (("Benchmarking and deployment. ", "b"), "Measure the accuracy cost of interpretability against gradient boosting, and package the pipeline as a triage tool for field teams."),
    ])

    # ------------------------------------------------------------ 9 conclusion
    rp.h("9. Conclusion", 1)
    rp.p(f"We asked whether a building's structure and location can predict its earthquake damage well enough to prioritise inspections. With a transparent, leakage-safe "
         f"logistic-regression pipeline, the answer is a qualified yes. The final model reaches a holdout micro-F1 of {f3(H_mic)} (macro-F1 {f3(H_mac)}) against a "
         f"{f3(base)} baseline, and its ranking finds {pct(R.cap(20))} of destroyed buildings in the first 20% of inspections. Location, encoded with cross-fitted "
         f"target encoding, contributed the most. Construction material contributed the most interpretable signal: mud-mortar stone raises risk and engineered RC lowers it. "
         f"An ordinal reformulation {'did not improve' if ord_mic <= 0.001 else 'improved'} on the multinomial model, and class weighting traded overall accuracy for "
         "recall of the rare class. The remaining errors are mostly between adjacent grades, and the largest further gains are likely to come from physical shaking data "
         "rather than from more complex models.")

    # ------------------------------------------------------------ references
    rp.h("References", 1)
    refs = [
        "DrivenData. (2019). Richter's Predictor: Modeling Earthquake Damage [Data set and competition]. https://www.drivendata.org/competitions/57/nepal-earthquake/",
        "Frank, E., & Hall, M. (2001). A simple approach to ordinal classification. In Proceedings of the 12th European Conference on Machine Learning (ECML), LNCS 2167, 145–156. Springer.",
        "Hastie, T., Tibshirani, R., & Friedman, J. (2009). The Elements of Statistical Learning (2nd ed.). Springer.",
        "Hosmer, D. W., Lemeshow, S., & Sturdivant, R. X. (2013). Applied Logistic Regression (3rd ed.). Wiley.",
        "Hunter, J. D. (2007). Matplotlib: A 2D graphics environment. Computing in Science & Engineering, 9(3), 90–95.",
        "McCullagh, P. (1980). Regression models for ordinal data. Journal of the Royal Statistical Society: Series B, 42(2), 109–142.",
        "McKinney, W. (2010). Data structures for statistical computing in Python. In Proceedings of the 9th Python in Science Conference, 56–61.",
        "Micci-Barreca, D. (2001). A preprocessing scheme for high-cardinality categorical attributes in classification and prediction problems. ACM SIGKDD Explorations, 3(1), 27–32.",
        "National Planning Commission. (2015). Nepal Earthquake 2015: Post Disaster Needs Assessment, Vol. A: Key Findings. Government of Nepal, Kathmandu.",
        "Pedregosa, F., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
    ]
    for ref in refs:
        para = doc.add_paragraph(ref)
        para.paragraph_format.left_indent = Cm(0.8)
        para.paragraph_format.first_line_indent = Cm(-0.8)
        para.paragraph_format.space_after = Pt(4)

    # ------------------------------------------------------------ appendix
    rp.page_break()
    rp.h("Appendix", 1)
    rp.h("A. Full cross-validation results", 2)
    rp.table(["ID", "Micro-F1", "Macro-F1", "Recall g1", "Recall g2", "Recall g3", "Max solver iter"],
             [[k, R.cv_cell(k), R.cv_cell(k, "f1_macro"), f3(R.rec(k, 1)), f3(R.rec(k, 2)), f3(R.rec(k, 3)),
               v["max_n_iter"]] for k, v in ex.items()],
             "All experiments: 5-fold CV mean ± sd; recall is the per-class CV mean; max solver iterations over the 5 folds.",
             widths=[1.1, 3.2, 3.2, 1.9, 1.9, 1.9, 2.4], font_size=8.5)
    rp.h("B. Regularisation path", 2)
    rp.table(["C", "Micro-F1 (mean)", "Micro-F1 (sd)", "Macro-F1 (mean)", "Max solver iterations"],
             [[f"{r['C']:g}", f4(r["mean"]), f4(r["std"]), f4(r["f1_macro_mean"]), int(r["max_n_iter"])] for r in trows],
             f"GridSearchCV results for {R.best}.", widths=[2.0, 3.2, 3.0, 3.2, 3.6])
    _appendix_figure(rp, "fig08_c_tuning", "B1", "CV micro-F1 as a function of C (log scale). The dashed line marks the selected value.", 11.5)
    rp.h("C. Additional exploration figures", 2)
    _appendix_figure(rp, "fig05_numeric_distributions", "C1", "Distributions of age, area, height (log scale) and floors, stacked by damage grade. Placeholder ages are excluded.")
    _appendix_figure(rp, "fig06_correlation", "C2", "Spearman correlations among numeric and selected material features and the damage grade.", 12.5)
    rp.h("D. Reproducibility", 2)
    m = R.r["meta"]
    rp.p(f"Environment: Python {m['python']}, scikit-learn {m['sklearn']} (Pedregosa et al., 2011), pandas {m['pandas']} (McKinney, 2010), NumPy {m['numpy']}, matplotlib {m['matplotlib']} (Hunter, 2007), seaborn {m['seaborn']}. "
         f"random_state = {m['random_state']} everywhere; {m['cv']}. The full pipeline runs with ", ("python -m src.run_all", "b"),
         f" (about {m.get('runtime_minutes', '–')} minutes for the notebook). Every number in this report is read from outputs/results.json, which the notebook writes. "
         f"The submission file outputs/submission.csv has {n(R.sub['n_rows'])} rows with columns building_id, damage_grade (integers 1–3). "
         f"Its predicted class shares are grade 1 {pct(R.sub['predicted_class_share']['1'])}, grade 2 {pct(R.sub['predicted_class_share']['2'])} "
         f"and grade 3 {pct(R.sub['predicted_class_share']['3'])}.")
    rp.h("E. Contribution statement", 2)
    rp.p(("[To be completed by the group: one line per member describing their contribution.]", "ph"))

    # ------------------------------------------------------------ page numbers
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE", "1")
    return rp


def _appendix_figure(rp, name, label, caption, width_cm=15):
    para = rp.doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.keep_with_next = True
    para.add_run().add_picture(str(FIG_DIR / f"{name}.png"), width=Cm(width_cm))
    c = rp.doc.add_paragraph(style="Caption")
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(f"Figure {label}. ")
    r.bold = True
    c.add_run(caption)


def export_pdf(docx_path=OUT_DOCX, pdf_path=OUT_PDF):
    """Update fields (TOC, page numbers) and export to PDF with Microsoft Word (Windows)."""
    if sys.platform != "win32":
        print("PDF export needs Microsoft Word on Windows; open report.docx and export manually.")
        return False
    try:
        import pythoncom
        import win32com.client
    except ImportError:
        print("pywin32 not installed; skipping PDF export.")
        return False
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        doc.Fields.Update()
        for i in range(1, doc.TablesOfContents.Count + 1):
            doc.TablesOfContents(i).Update()
        doc.Save()
        doc.SaveAs2(str(pdf_path.resolve()), FileFormat=17)  # wdFormatPDF
        pages = doc.ComputeStatistics(2)  # wdStatisticPages
        doc.Close(False)
        print(f"wrote {pdf_path} ({pages} pages)")
        return True
    finally:
        word.Quit()


def main():
    R = ResultsView()
    rp = build(R)
    OUT_DOCX.parent.mkdir(exist_ok=True)
    rp.doc.save(OUT_DOCX)
    print(f"wrote {OUT_DOCX} ({rp.fig_no} numbered figures, {rp.tab_no} tables)")
    export_pdf()


if __name__ == "__main__":
    main()
